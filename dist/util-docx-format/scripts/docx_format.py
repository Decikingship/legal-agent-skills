#!/usr/bin/env python3
"""
法律文书 docx 排版

把 Markdown 或纯文本的法律文书排成符合提交规范的 Word 文档。

**中文法律文书排版的坑集中在几处，全部与中文字体设置有关：**
  1. python-docx 设 `font.name` 只作用于西文，**中文必须另设 `w:eastAsia`**，
     否则中文仍是默认字体，打印出来不是要求的字体
  2. 行距用「固定值」而非「多倍行距」——法院与仲裁委的格式要求通常是固定磅值
  3. 首行缩进按**字符**（2 字符）而非厘米，字号变化时缩进要跟着变

预置三套规范：court（法院文书）/ arbitration（仲裁文书）/ opinion（法律意见书）。

用法：
    docx_format.py --input 起诉状.md --output 起诉状.docx --style court \
        --title "民事起诉状"
"""

import argparse
import os
import re
import sys

try:
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
    from docx.oxml.ns import qn
    from docx.shared import Cm, Pt
except ImportError:
    sys.exit("错误：需要 python-docx。安装：pip install python-docx --break-system-packages")

# 三套预置规范。字号用磅值：小二18 / 三号16 / 小三15 / 四号14 / 小四12 / 五号10.5
STYLES = {
    "court": {
        "名称": "法院诉讼文书",
        "页边距": {"top": 3.7, "bottom": 3.5, "left": 2.8, "right": 2.6},
        "标题": {"font": "方正小标宋简体", "fallback": "宋体", "size": 22, "bold": False,
                 "align": "center", "space_after": 24},
        "正文": {"font": "仿宋_GB2312", "fallback": "仿宋", "size": 16,
                 "line": 28, "indent_chars": 2},
        "一级标题": {"font": "黑体", "size": 16, "bold": False},
        "落款": {"align": "right"},
        "说明": "参照公文格式，正文仿宋三号、固定行距28磅。**各地法院立案窗口可能另有要求，提交前确认**",
    },
    "arbitration": {
        "名称": "仲裁文书",
        "页边距": {"top": 3.0, "bottom": 3.0, "left": 3.0, "right": 2.5},
        "标题": {"font": "黑体", "fallback": "黑体", "size": 18, "bold": True,
                 "align": "center", "space_after": 20},
        "正文": {"font": "宋体", "fallback": "宋体", "size": 12,
                 "line": 24, "indent_chars": 2},
        "一级标题": {"font": "黑体", "size": 14, "bold": True},
        "落款": {"align": "right"},
        "说明": "仲裁机构对格式要求通常宽于法院，**以受理机构的规则为准**",
    },
    "opinion": {
        "名称": "法律意见书",
        "页边距": {"top": 2.54, "bottom": 2.54, "left": 3.17, "right": 3.17},
        "标题": {"font": "黑体", "fallback": "黑体", "size": 18, "bold": True,
                 "align": "center", "space_after": 18},
        "正文": {"font": "宋体", "fallback": "宋体", "size": 12,
                 "line": 22, "indent_chars": 2},
        "一级标题": {"font": "黑体", "size": 14, "bold": True},
        "落款": {"align": "right"},
        "说明": "对外出具的法律意见书，可按律所模板调整",
    },
}


def set_font(run, name, size, bold=False):
    """设置字体。**中文必须单独设 eastAsia**，只设 font.name 中文不生效。"""
    run.font.name = name
    run.font.size = Pt(size)
    run.font.bold = bold
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)


def set_para(par, cfg, align=None, indent=True):
    pf = par.paragraph_format
    if cfg.get("line"):
        # 固定行距，不用多倍行距——法院格式要求通常是固定磅值
        pf.line_spacing_rule = WD_LINE_SPACING.EXACTLY
        pf.line_spacing = Pt(cfg["line"])
    if indent and cfg.get("indent_chars"):
        # 首行缩进按字符：字号 × 字符数，字号变则缩进跟着变
        pf.first_line_indent = Pt(cfg["size"] * cfg["indent_chars"])
    if cfg.get("space_after"):
        pf.space_after = Pt(cfg["space_after"])
    amap = {"center": WD_ALIGN_PARAGRAPH.CENTER, "right": WD_ALIGN_PARAGRAPH.RIGHT,
            "left": WD_ALIGN_PARAGRAPH.LEFT}
    if align in amap:
        par.alignment = amap[align]
    elif not align:
        par.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY


def build(text, style, title, out):
    st = STYLES[style]
    doc = Document()
    sec = doc.sections[0]
    m = st["页边距"]
    sec.top_margin, sec.bottom_margin = Cm(m["top"]), Cm(m["bottom"])
    sec.left_margin, sec.right_margin = Cm(m["left"]), Cm(m["right"])

    warnings = []
    if title:
        p = doc.add_paragraph()
        set_para(p, st["标题"], st["标题"]["align"], indent=False)
        set_font(p.add_run(title), st["标题"]["font"], st["标题"]["size"],
                 st["标题"].get("bold", False))

    for raw in text.split("\n"):
        line = raw.rstrip()
        if not line.strip():
            continue
        # Markdown 标题 → 一级标题样式
        m2 = re.match(r"^(#{1,4})\s+(.*)$", line)
        if m2:
            p = doc.add_paragraph()
            set_para(p, {"size": st["一级标题"]["size"], "line": st["正文"]["line"]},
                     "left", indent=False)
            set_font(p.add_run(m2.group(2)), st["一级标题"]["font"],
                     st["一级标题"]["size"], st["一级标题"].get("bold", False))
            continue
        # 中文序号开头的段落：可能是**章节标题**，也可能是**条目正文**。
        # 教训：初版一律当标题，结果诉讼请求的「一、判令被告支付货款…」
        # 被顶格加粗成了标题——而它是条目正文。
        # 区分依据：标题短且不以标点结尾；条目长且以「；」「。」收尾。
        m3 = re.match(r"^\s*(?:[一二三四五六七八九十]+\s*[、.．]|"
                      r"（[一二三四五六七八九十]+）)\s*(.*)$", line)
        if m3:
            body = m3.group(1).strip()
            is_heading = len(body) <= 20 and not re.search(r"[；。，、]$", body)
            p = doc.add_paragraph()
            if is_heading:
                set_para(p, st["正文"], "left", indent=False)
                set_font(p.add_run(line.strip()), st["一级标题"]["font"],
                         st["正文"]["size"], True)
            else:
                # 条目正文：与正文同字体，首行缩进，不加粗
                set_para(p, st["正文"])
                set_font(p.add_run(re.sub(r"\*\*(.+?)\*\*", r"\1", line.strip())),
                         st["正文"]["font"], st["正文"]["size"])
            continue
        # 落款（此致/具状人/年月日）右对齐
        if re.match(r"^\s*(此致|具状人|申请人|代理人|答辩人|上诉人)[:：]?\s*$", line) or \
           re.match(r"^\s*\d{4}\s*年\s*\d{1,2}\s*月\s*\d{1,2}\s*日\s*$", line):
            p = doc.add_paragraph()
            set_para(p, st["正文"], st["落款"]["align"], indent=False)
            set_font(p.add_run(line.strip()), st["正文"]["font"], st["正文"]["size"])
            continue
        # Markdown 表格与强调标记：法律文书不用，剥离后提示
        if line.lstrip().startswith("|"):
            warnings.append("检测到 Markdown 表格。**正式法律文书中的表格须手工排版**，"
                            "本脚本已按纯文本处理")
        clean = re.sub(r"\*\*(.+?)\*\*", r"\1", line)
        clean = re.sub(r"\[([^\]]+)\]\([^)]+\)", r"\1", clean)   # 链接只留文字
        if clean != line:
            warnings.append("已剥离 Markdown 强调与链接标记——"
                            "**正式文书中法条与案例的出处应写在正文里，不用超链接**")
        p = doc.add_paragraph()
        set_para(p, st["正文"])
        set_font(p.add_run(clean.strip()), st["正文"]["font"], st["正文"]["size"])

    doc.save(out)
    return sorted(set(warnings)), st


def main():
    p = argparse.ArgumentParser(description="法律文书 docx 排版",
                                formatter_class=argparse.RawDescriptionHelpFormatter)
    p.add_argument("--input", required=True, help="输入 md/txt 文件")
    p.add_argument("--output", required=True, help="输出 docx 路径")
    p.add_argument("--style", default="court", choices=list(STYLES))
    p.add_argument("--title", help="文书标题，如「民事起诉状」")
    args = p.parse_args()

    if not os.path.exists(args.input):
        sys.exit(f"错误：输入文件不存在 {args.input}")
    text = open(args.input, encoding="utf-8").read()
    warnings, st = build(text, args.style, args.title, args.output)

    import json
    print(json.dumps({
        "输出": args.output, "规范": st["名称"],
        "页边距cm": st["页边距"],
        "正文": f"{st['正文']['font']} {st['正文']['size']}磅，固定行距{st['正文']['line']}磅，"
                f"首行缩进{st['正文']['indent_chars']}字符",
        "警告": warnings or "无",
        "提示": [
            st["说明"],
            "**中文字体必须设 eastAsia**，只设 font.name 中文不生效——"
            "这是 python-docx 排中文最常见的坑，打印出来才发现字体不对",
            "生成后**务必转 PDF 目视检查**："
            "`soffice --headless --convert-to pdf out.docx`，再看页码、分页、字体是否正确",
            "⚠️ **各地法院立案窗口的格式要求可能不同**，"
            "提交前确认受理法院有无特别要求（份数、装订、纸张）",
            "表格、附件清单、证据目录建议在 Word 中手工排版，脚本只处理正文",
        ],
    }, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
